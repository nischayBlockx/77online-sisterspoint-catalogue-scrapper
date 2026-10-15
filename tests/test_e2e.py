"""End-to-end: run the real bot against the local mock B2B site."""
import csv
import os

import pytest
from openpyxl import load_workbook

from sp_scraper.cli import main
from tests.mock_site.server import PASSWORD, USER, start

pytest.importorskip("playwright")


@pytest.fixture()
def site(tmp_path, monkeypatch):
    srv, url = start()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("SP_BASE_URL", url)
    monkeypatch.setenv("SP_USERNAME", USER)
    monkeypatch.setenv("SP_PASSWORD", PASSWORD)
    (tmp_path / "config.yaml").write_text("delay_seconds: 0\ntimeout_seconds: 20\n")
    yield url
    srv.shutdown()


def test_wrong_password(site, monkeypatch):
    monkeypatch.setenv("SP_PASSWORD", "nope")
    assert main(["login"]) == 2


def test_full_scrape_by_category(site, tmp_path):
    assert main(["scrape"]) == 0
    out = tmp_path / "output"

    wb = load_workbook(out / "catalogue_all.xlsx")
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 9  # every size of every colour
    by_ean = {str(r["EAN"]): r for r in rows}

    grid = by_ean["5712345000011"]  # size-grid product
    assert (grid["StyleCode"], grid["StyleName"], grid["Color"], grid["Size"]) == ("SP100", "Gerda Dress", "Black", "S")
    assert (grid["Wholesale_Price_Euro"], grid["Retail_Price_Euro"], grid["Available Stock"]) == ("24,95", "59,95", 12)
    assert (grid["ProductsComposition"], grid["OriginCountry"], grid["CustomTariffNo"], grid["Season"]) == ("100% Viscose", "IN", "62044400", "FW26")
    assert grid["ItemMainCategory"] == "Dresses"
    assert grid["ImageProductP1"].endswith("/img/SP100-black_1.jpg")
    assert grid["ImageProductP2"].endswith("/img/SP100-black_2.jpg")  # lazy-loaded image
    assert "logo" not in " ".join(str(v) for v in grid.values())

    ld = by_ean["5712345000073"]  # JSON-LD product (page 2 of the category)
    assert (ld["Brand"], ld["Size"], ld["ItemCode"], ld["Available Stock"], ld["Wholesale_Price_Euro"]) == ("Sisters Point", "S", "SP200-S", 2, "29,95")

    xhr = by_ean["5712345000097"]  # sizes/EANs only available via XHR JSON, JS-rendered listing
    assert (xhr["ItemMainCategory"], xhr["Size"], xhr["Available Stock"], xhr["ItemCode"]) == ("Tops", "L", 5, "SP300-white-L")

    assert sorted(p.name for p in (out / "by_category").iterdir()) == [
        "catalogue_dresses.xlsx", "catalogue_tops.xlsx", "shopify_dresses.csv", "shopify_tops.csv"]

    with open(out / "shopify_products_all.csv", newline="", encoding="utf-8") as fh:
        shop = list(csv.DictReader(fh))
    gerda = [r for r in shop if r["Handle"] == "gerda-dress-sp100"]
    assert gerda[0]["Title"] == "Gerda Dress" and gerda[0]["Option1 Name"] == "Color" and gerda[0]["Option2 Name"] == "Size"
    assert {(r["Option1 Value"], r["Option2 Value"]) for r in gerda if r["Variant Barcode"]} == {
        ("Black", "S"), ("Black", "M"), ("Black", "L"), ("Blue", "S"), ("Blue", "M")}
    assert gerda[0]["Variant Price"] == "59.95" and gerda[0]["Cost per item"] == "24.95"

    # resume: a second run scrapes nothing new and output stays identical
    assert main(["scrape"]) == 0
    assert load_workbook(out / "catalogue_all.xlsx").active.max_row == 10


def test_single_category_and_limit(site, tmp_path):
    assert main(["scrape", "--category", "tops", "--limit", "1"]) == 0
    rows = list(load_workbook(tmp_path / "output" / "catalogue_all.xlsx").active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2 and all(r[1] == "Tops" for r in rows)


@pytest.mark.skipif(not os.environ.get("CI_FULL"), reason="debug dump smoke test")
def test_discover(site, tmp_path):
    assert main(["discover"]) == 0
    assert (tmp_path / "output" / "debug" / "04_product.png").exists()
