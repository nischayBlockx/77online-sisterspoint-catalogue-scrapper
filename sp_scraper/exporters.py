"""Exporters: scraped products -> reference catalogue rows -> Excel / Shopify CSV."""
from __future__ import annotations

import csv
import datetime as dt
import json
import re
import unicodedata
from collections import OrderedDict
from dataclasses import asdict
from pathlib import Path
from typing import Iterable

from .extract import parse_price
from .models import CATALOGUE_COLUMNS, IMAGE_COLUMNS, Product, Variant


# --------------------------------------------------------------------------- formatting
def fmt_decimal(value: float | None, decimal_comma: bool = True) -> str:
    if value is None:
        return ""
    s = f"{value:.2f}".rstrip("0").rstrip(".")
    return s.replace(".", ",") if decimal_comma else s


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


# --------------------------------------------------------------------------- products -> rows
def base_style_code(code: str) -> str:
    """Style code without the colour suffix (e.g. '19059-904' -> '19059')."""
    m = re.fullmatch(r"(\d+)-\d+", (code or "").strip())
    return m.group(1) if m else (code or "")


def product_to_rows(p: Product, cfg: dict | None = None) -> list[dict]:
    cfg = cfg or {}
    decimal_comma = cfg.get("decimal_comma", True)
    today = dt.date.today().isoformat()
    skus = [v.sku for v in p.variants if v.sku]
    has_dup_skus = len(skus) != len(set(skus))
    rows = []
    for v in p.variants:
        images = v.images or p.images
        row = {c: "" for c in CATALOGUE_COLUMNS}
        size = v.size or "one size"
        color = v.color or p.color
        row.update({
            "Brand": p.brand or cfg.get("default_brand", ""),
            "ItemMainCategory": p.category,
            "EAN": v.ean,
            "ItemCode": (v.sku if v.sku and not has_dup_skus else "-".join(x for x in (p.style_code, slugify(color), slugify(size)) if x)),
            "ItemDescription": " ".join(x for x in (p.style_name, color, size) if x),
            "Color": color,
            "Size": size,
            "StyleCode": base_style_code(p.style_code),
            "StyleName": p.style_name,
            "Wholesale_Price_Euro": fmt_decimal(v.wholesale_price if v.wholesale_price is not None else p.wholesale_price, decimal_comma),
            "Retail_Price_Euro": fmt_decimal(v.retail_price if v.retail_price is not None else p.retail_price, decimal_comma),
            "Weight": p.weight.replace(".", ",") if decimal_comma else p.weight,
            "WeightUnit": p.weight_unit or ("kg" if p.weight else ""),
            "Season": p.season,
            "CustomTariffNo": re.sub(r"\s", "", p.tariff),
            "Gender": p.gender or cfg.get("default_gender", ""),
            "OriginCountry": p.origin,
            "ProductsComposition": p.composition,
            "CareInstruction": p.care,
            "ProductText": p.description,
            "CreationDate": today,
            "Available Stock": "" if v.stock is None else v.stock,
        })
        for col, img in zip(IMAGE_COLUMNS, images):
            row[col] = img
        rows.append(row)
    return rows


def products_to_rows(products: Iterable[Product], cfg: dict | None = None) -> list[dict]:
    rows, seen = [], set()
    for p in products:
        for r in product_to_rows(p, cfg):
            key = r["EAN"] or (r["ItemCode"], r["Color"], r["Size"])
            if key in seen:
                continue
            seen.add(key)
            rows.append(r)
    return rows


def load_products_jsonl(path: Path) -> list[Product]:
    products = []
    if not path.exists():
        return products
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            if not line.strip():
                continue
            d = json.loads(line)
            d["variants"] = [Variant(**v) for v in d.get("variants", [])]
            products.append(Product(**d))
    return products


def append_product_jsonl(path: Path, p: Product) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(asdict(p), ensure_ascii=False) + "\n")


# --------------------------------------------------------------------------- Excel (reference format)
def write_catalogue_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(CATALOGUE_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        values = []
        for c in CATALOGUE_COLUMNS:
            v = r.get(c, "")
            if c == "EAN" and str(v).isdigit():
                v = int(v)  # reference stores EANs as numbers
            values.append(v)
        ws.append(values)
    ean_col = CATALOGUE_COLUMNS.index("EAN") + 1
    for (cell,) in ws.iter_rows(min_row=2, min_col=ean_col, max_col=ean_col):
        cell.number_format = "0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(CATALOGUE_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(c) + 4))
    wb.save(path)


def read_catalogue_xlsx(path: Path) -> list[dict]:
    """Read a catalogue in the reference format (e.g. the TBI workbook)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = []
    for values in it:
        if not any(v not in (None, "") for v in values):
            continue
        row = {}
        for h, v in zip(header, values):
            if v is None:
                v = ""
            elif isinstance(v, float) and v.is_integer() and h in ("EAN", "Available Stock"):
                v = int(v)
            elif isinstance(v, (dt.date, dt.datetime)):
                v = v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
            row[h] = v
        rows.append(row)
    return rows


# --------------------------------------------------------------------------- Shopify CSV
SHOPIFY_COLUMNS = [
    "Handle", "Title", "Body (HTML)", "Vendor", "Product Category", "Type", "Tags", "Published",
    "Option1 Name", "Option1 Value", "Option2 Name", "Option2 Value", "Option3 Name", "Option3 Value",
    "Variant SKU", "Variant Grams", "Variant Inventory Tracker", "Variant Inventory Qty",
    "Variant Inventory Policy", "Variant Fulfillment Service", "Variant Price", "Variant Compare At Price",
    "Variant Requires Shipping", "Variant Taxable", "Variant Barcode", "Image Src", "Image Position",
    "Image Alt Text", "Gift Card", "SEO Title", "SEO Description", "Variant Image", "Variant Weight Unit",
    "Cost per item", "Status",
]

SIZE_ORDER = ["XXXS", "XXS", "XS", "S", "M", "L", "XL", "XXL", "2XL", "XXXL", "3XL", "4XL", "5XL", "6XL"]


def _size_key(size: str):
    s = str(size).upper().strip()
    if s in SIZE_ORDER:
        return (0, SIZE_ORDER.index(s), s)
    m = re.match(r"(\d+(?:[.,]\d+)?)", s)
    if m:
        return (1, float(m.group(1).replace(",", ".")), s)
    return (2, 0, s)


def _num(v) -> float | None:
    if v in (None, ""):
        return None
    return parse_price(v)


def _grams(weight, unit) -> str:
    w = _num(weight)
    if w is None:
        return ""
    unit = (str(unit) or "kg").lower()
    factor = {"kg": 1000, "g": 1, "gr": 1, "lb": 453.592, "lbs": 453.592, "oz": 28.3495}.get(unit, 1000)
    return str(round(w * factor))


def _body_html(row: dict) -> str:
    from html import escape

    parts = []
    if row.get("ProductText"):
        parts.append(f"<p>{escape(str(row['ProductText']))}</p>")
    details = [
        ("Composition", row.get("ProductsComposition")),
        ("Care", row.get("CareInstruction")),
        ("Country of origin", row.get("OriginCountry")),
    ]
    items = "".join(f"<li><strong>{k}:</strong> {escape(str(v))}</li>" for k, v in details if v)
    if items:
        parts.append(f"<ul>{items}</ul>")
    return "".join(parts)


def rows_to_shopify(rows: list[dict], cfg: dict | None = None) -> list[dict]:
    """Group catalogue rows by StyleCode into Shopify products with Color/Size options."""
    cfg = cfg or {}
    shop = cfg.get("shopify", {})
    multiplier = float(shop.get("price_multiplier", 1.0))
    markup = float(shop.get("markup_on_wholesale", 2.5))
    status = shop.get("status", "draft")
    published = "TRUE" if shop.get("published", False) else "FALSE"
    inventory = bool(shop.get("track_inventory", True))
    vendor_override = shop.get("vendor") or ""
    extra_tags = list(shop.get("tags", []))
    compare_at_retail = bool(shop.get("compare_at_retail", False))

    groups: "OrderedDict[str, list[dict]]" = OrderedDict()
    for r in rows:
        key = str(r.get("StyleCode") or r.get("StyleName") or r.get("ItemCode"))
        groups.setdefault(key, []).append(r)

    out: list[dict] = []
    used_handles: set[str] = set()
    for style, items in groups.items():
        first = items[0]
        title = str(first.get("StyleName") or first.get("ItemDescription") or style)
        handle = slugify(f"{title}-{style}") if style and slugify(style) not in slugify(title) else slugify(title)
        base_handle, n = handle, 2
        while handle in used_handles:
            handle, n = f"{base_handle}-{n}", n + 1
        used_handles.add(handle)

        colors = list(OrderedDict.fromkeys(str(i.get("Color") or "") for i in items))
        has_color = any(colors)
        has_size = any(str(i.get("Size") or "").lower() not in ("", "one size", "onesize", "os") for i in items) or not has_color
        items = sorted(items, key=lambda i: (colors.index(str(i.get("Color") or "")), _size_key(i.get("Size") or "")))

        # product images: all images of all colours, de-duplicated, in column order
        images: list[str] = []
        color_image: dict[str, str] = {}
        for i in items:
            for col in IMAGE_COLUMNS:
                u = str(i.get(col) or "").strip()
                if u:
                    color_image.setdefault(str(i.get("Color") or ""), u)
                    if u not in images:
                        images.append(u)

        tags = [t for t in [first.get("Brand"), first.get("ItemMainCategory"), first.get("Season"), first.get("Gender")] if t]
        tags += extra_tags
        seen_option_values: set[tuple] = set()
        variant_rows = []
        for idx, i in enumerate(items):
            opt1 = str(i.get("Color") or "") if has_color else ""
            opt2 = str(i.get("Size") or "one size") if has_size else ""
            if not has_color and not has_size:
                opt1 = "Default Title"
            if (opt1, opt2) in seen_option_values:
                continue  # Shopify rejects duplicate option combinations
            seen_option_values.add((opt1, opt2))
            wholesale = _num(i.get("Wholesale_Price_Euro"))
            retail = _num(i.get("Retail_Price_Euro"))
            price = retail if retail else (wholesale * markup if wholesale else None)
            price = round(price * multiplier, 2) if price else None
            cost = round(wholesale * multiplier, 2) if wholesale else None
            stock = i.get("Available Stock")
            row = {c: "" for c in SHOPIFY_COLUMNS}
            row.update({
                "Handle": handle,
                "Option1 Name": ("Color" if has_color else ("Size" if has_size else "Title")) if not variant_rows else "",
                "Option1 Value": opt1 if (has_color or not has_size) else opt2,
                "Option2 Name": ("Size" if has_color and has_size else "") if not variant_rows else "",
                "Option2 Value": opt2 if has_color and has_size else "",
                "Variant SKU": i.get("ItemCode") or "",
                "Variant Grams": _grams(i.get("Weight"), i.get("WeightUnit")),
                "Variant Inventory Tracker": "shopify" if inventory else "",
                "Variant Inventory Qty": stock if inventory and stock not in (None, "") else "",
                "Variant Inventory Policy": "deny",
                "Variant Fulfillment Service": "manual",
                "Variant Price": f"{price:.2f}" if price else "",
                "Variant Compare At Price": "",
                "Variant Requires Shipping": "TRUE",
                "Variant Taxable": "TRUE",
                "Variant Barcode": str(i.get("EAN") or ""),
                "Variant Image": color_image.get(str(i.get("Color") or ""), ""),
                "Variant Weight Unit": "kg",
                "Cost per item": f"{cost:.2f}" if cost else "",
            })
            if compare_at_retail and retail and wholesale and price and retail * multiplier > price:
                row["Variant Compare At Price"] = f"{retail * multiplier:.2f}"
            variant_rows.append(row)

        # product-level fields go on the first row only
        variant_rows[0].update({
            "Title": title,
            "Body (HTML)": _body_html(first),
            "Vendor": vendor_override or first.get("Brand") or "",
            "Product Category": shop.get("product_category", ""),
            "Type": first.get("ItemMainCategory") or "",
            "Tags": ", ".join(dict.fromkeys(str(t) for t in tags)),
            "Published": published,
            "Gift Card": "FALSE",
            "SEO Title": title,
            "SEO Description": str(first.get("ProductText") or "")[:320],
            "Status": status,
        })
        if len(variant_rows) > 2048:
            print(f"warning: {handle} has {len(variant_rows)} variants (Shopify max 2048)")

        # attach images: one per row; extra image-only rows after the variants
        for pos, url in enumerate(images, 1):
            if pos <= len(variant_rows):
                target = variant_rows[pos - 1]
            else:
                target = {c: "" for c in SHOPIFY_COLUMNS}
                target["Handle"] = handle
                variant_rows.append(target)
            target["Image Src"] = url
            target["Image Position"] = str(pos)
            target["Image Alt Text"] = title
        out.extend(variant_rows)
    return out


def write_shopify_csv(rows: list[dict], path: Path, cfg: dict | None = None) -> int:
    shop_rows = rows_to_shopify(rows, cfg)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=SHOPIFY_COLUMNS)
        w.writeheader()
        w.writerows(shop_rows)
    return len({r["Handle"] for r in shop_rows})
